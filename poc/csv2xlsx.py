#!/usr/bin/env python3
"""
csv2xlsx.py — wandelt Benchmark-CSV-Dateien in eine .xlsx-Datei um.

Verwendet openpyxl (im venv installierbar: pip install openpyxl).

Sheets:
  1. "Rohdaten"   — jede einzelne Messung (fuer eigene statistische Auswertungen)
  2. "Auswertung" — Kennzahlen pro Algorithmus/Operation
                    (n, Mittelwert, Standardabweichung, Median, Min, P95, Max; in us)
  3. "Meta"       — optional: Metadaten pro Algorithmus aus einer Meta-CSV
                    (z.B. dnssec_bench_meta.csv: Schluessel-/Signaturgroessen)

Aufruf:
  python3 csv2xlsx.py [eingabe.csv] [ausgabe.xlsx] [meta.csv]
  (Standard: pqc_bench.csv -> pqc_bench.xlsx, ohne Meta)
"""

import csv
import math
import statistics
import sys

from openpyxl import Workbook
from openpyxl.styles import Font


def read_measurements(csv_path):
    measurements = []
    with open(csv_path, newline="") as f:
        for row in csv.DictReader(f):
            measurements.append((
                row["algorithm"],
                row["operation"],
                int(row["iteration"]),
                int(row["time_ns"]),
            ))
    return measurements


def main():
    csv_path = sys.argv[1] if len(sys.argv) > 1 else "pqc_bench.csv"
    xlsx_path = sys.argv[2] if len(sys.argv) > 2 else "pqc_bench.xlsx"
    meta_path = sys.argv[3] if len(sys.argv) > 3 else None

    measurements = read_measurements(csv_path)
    if not measurements:
        sys.exit(f"Keine Daten in {csv_path}")

    wb = Workbook()
    bold = Font(bold=True)

    # --- Sheet 1: Rohdaten ---
    ws = wb.active
    ws.title = "Rohdaten"
    ws.append(["algorithm", "operation", "iteration", "time_ns", "time_us"])
    for cell in ws[1]:
        cell.font = bold
    for alg, op, it, ns in measurements:
        ws.append([alg, op, it, ns, round(ns / 1000.0, 3)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"

    # --- Sheet 2: Auswertung ---
    groups = {}
    for alg, op, _it, ns in measurements:
        groups.setdefault((alg, op), []).append(ns)

    ws = wb.create_sheet("Auswertung")
    ws.append(["algorithm", "operation", "n",
               "mean_us", "stddev_us", "median_us", "min_us", "p95_us", "max_us"])
    for cell in ws[1]:
        cell.font = bold
    for (alg, op), values in groups.items():
        n = len(values)
        us = [v / 1000.0 for v in values]
        ordered = sorted(us)
        p95 = ordered[max(0, math.ceil(0.95 * n) - 1)]
        ws.append([
            alg, op, n,
            round(statistics.fmean(us), 3),
            round(statistics.stdev(us), 3) if n > 1 else 0.0,
            round(statistics.median(us), 3),
            round(ordered[0], 3),
            round(p95, 3),
            round(ordered[-1], 3),
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    # --- Sheet 3: Meta (optional) ---
    if meta_path is not None:
        ws = wb.create_sheet("Meta")
        with open(meta_path, newline="") as f:
            for i, row in enumerate(csv.reader(f)):
                ws.append([_maybe_number(v) for v in row])
                if i == 0:
                    for cell in ws[1]:
                        cell.font = bold
        ws.freeze_panes = "A2"

    for sheet in wb.worksheets:
        for column in sheet.columns:
            width = max(len(str(c.value)) for c in column if c.value is not None)
            sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 40)

    wb.save(xlsx_path)
    print(f"{xlsx_path}: {len(measurements)} Messungen, "
          f"{len(groups)} Algorithmus/Operation-Kombinationen")


def _maybe_number(value):
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


if __name__ == "__main__":
    main()
